from pathlib import Path
import re, json, copy
from urllib.parse import quote, unquote, urlparse
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(r'D:/Evo Web Design/FarmGame')
SOURCE = Path(r'C:/Users/esinencuP/Downloads/My_Little_Farm_Attestation_Report.docx')
OUT = ROOT / 'Docs/Attestation/My_Little_Farm_Attestation_Report_Expanded.docx'
WORK = Path(r'C:/Users/esinencuP/.codex/visualizations/2026/10/12/01a1316d-41a6-7b00-8409-a5110ef7f494/attestation_work')
WORK.mkdir(parents=True, exist_ok=True)
doc = Document(SOURCE)
orig = list(doc.paragraphs)
original_heading_names = [p.text for p in orig if p.style.name.startswith('Heading')]
records = []

def link(p, label, target=None, anchor=None, size=9.5):
    h = OxmlElement('w:hyperlink')
    if target:
        rid = p.part.relate_to(target, RT.HYPERLINK, is_external=True)
        h.set(qn('r:id'), rid)
    if anchor: h.set(qn('w:anchor'), anchor)
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'),'1F4E79'); pr.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'),'single'); pr.append(u)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'),str(int(size*2))); pr.append(sz)
    r.append(pr); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)

def resolve(path):
    p = Path(path)
    p = p if p.is_absolute() else ROOT / p
    if not p.exists(): raise FileNotFoundError(str(p))
    return p

def source_ref(p, path, token='FILE', count=0):
    f=resolve(path)
    p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.keep_with_next=False
    display = f.relative_to(ROOT).as_posix() if f.is_relative_to(ROOT) else f.name
    if token=='FILE':
        link(p,display,f.as_uri())
        records.append({'path':str(f),'start':None,'end':None,'token':token})
        return
    lines=f.read_text(encoding='utf-8-sig').splitlines()
    hits=[i for i,l in enumerate(lines) if token in l]
    if not hits: raise ValueError(f'Missing source token: {path} :: {token}')
    start=hits[0]+1; end=min(len(lines), start+int(count)-1)
    link(p,display,f.as_uri())
    r=p.add_run(f' · строки {start}–{end} · {token.strip()} ')
    r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('444444')
    uri='vscode://file/'+quote(f.as_posix(),safe='/:')+f':{start}'
    link(p,'к строке',uri,size=9)
    records.append({'path':str(f),'start':start,'end':end,'token':token})

def before(anchor,text='',style='Normal'):
    return anchor.insert_paragraph_before(text,style=style)

replacements={
4:'Аттестация № 1 · расширенная версия · 12 октября 2026 года',
13:'Приложения: полный перечень 86 C#-файлов, конфигурация проекта, отчёты проверок и источники с активными ссылками.',
14:'Ссылки [П1–П10] обозначают материалы проекта, [И1–И15] — официальную документацию. Технические ссылки ведут к файлам полного проекта; рядом указаны методы и номера строк. Навигация и сведения об открытии ссылок приведены далее.',
20:'Состав полного проекта',
22:'Unity 6000.6.4f1 подтверждён ProjectSettings/ProjectVersion.txt. В Packages/manifest.json указаны Input System 1.20.0, Test Framework 1.8.0 и uGUI 2.6.0. GraphicsSettings использует Built-in Render Pipeline. Эти сведения проверяются по конфигурации полного проекта, а не только по UnityProjectContext [П6, П8].',
23:'Проверены исходники, структура сборок, сериализованная сцена, Animator Controller, определения контента, конфигурация и отчёты прошлых запусков. 12 октября повторно выполнен Verify-Scripts.ps1: синтаксис и сборки 86 файлов, чистые контракты сетки, экономики и строительства прошли. Unity Test Runner и новый build в этом проходе не запускались; результаты Unity от 8 октября приводятся отдельно с их датой [П9].',
24:'В полном проекте доступны Assets, Packages и ProjectSettings, импортированные модели, материалы, клипы, авторские карты и результаты визуальной проверки. Стартовая сцена — Assets/_Game/Scenes/Prototype.unity. Для защиты можно проследить связь между игровым компонентом, контентным asset и реальным объектом в Hierarchy.',
92:'EditMode-проверки охватывают арифметику, координаты, footprint и расчёт стадий. PlayMode-сценарии проверяют игровой цикл, деревья, кур, сохранения и регрессии выбора цели. MvpReliabilityTests проверяет backup, отключение HUD и плавность. В отчёте последнего прохода от 8 октября зафиксированы 11/11 EditMode и 25/25 PlayMode; это сохранённый результат, а не новый запуск 12 октября [П9, И9].',
131:'В проекте подготовлены и ранее выполнены EditMode и PlayMode тесты. Они проверяют расчёт стадий, экономику, размещение и восстановление, а также выявленные ошибки взаимодействия. В отчёте от 8 октября зафиксированы 11 EditMode и 25 PlayMode без отказов. Отдельная повторная проверка текущих исходников выполнена 12 октября. Я могу показать конкретный сценарий и объяснить, какое поведение он контролирует; для актуального состояния Test Runner нужен отдельный запуск.',
261:'Критерий следующего этапа: игрок без помощи проходит основной цикл; одна зона визуально согласована; анимации и коллизии работают; сохранение восстанавливает результат; вклад обоих участников можно предъявить. Корова и переработка остаются задачами плана, а не подтверждёнными функциями текущего проекта.',
277:'Пути приведены относительно Assets/_Game в полном Unity-проекте. Названия файлов в таблицах являются активными ссылками на локальные исходники. Краткое назначение не означает, что каждый файл является отдельной законченной механикой.',
285:'П1 Полный Unity проект FarmGame',
286:'Первичный источник текущей реализации: D:/Evo Web Design/FarmGame. Проверены C#-скрипты, сцена Prototype, Animator Controller, определения контента, настройки проекта, зависимости, материалы и импортированные ассеты.',
289:'П3 My Little Farm GDD',
291:'П4 План разработки и дедлайны',
293:'П5 Unity Guide',
304:'Официальные источники объясняют используемые API и принципы. Большинство сохранённых ссылок относится к Unity 6.0, источник по тестам — к Unity 6.4. Input System документация приведена для API 1.20. Версии самого проекта подтверждаются ProjectVersion и manifest, а не версией страницы документации. Основные API движения, Blend Tree, JSON и File.Replace дополнительно сверены 12 октября 2026 года.'
}
for i,text in replacements.items(): orig[i].text=text

# Remove old manual page-break-only paragraphs; the original section ladder remains intact.
for p in orig:
    for b in list(p._p.xpath('.//w:br')):
        if b.get(qn('w:type'))=='page': b.getparent().remove(b)
    if p.style.name=='Heading 1':
        p.paragraph_format.page_break_before=False
        p.paragraph_format.keep_with_next=True
orig[1].paragraph_format.page_break_before=False
orig[16].paragraph_format.page_break_before=True

for name,size in [('Normal',11),('Title',26),('Subtitle',12),('Heading 1',18),('Heading 2',13),('Heading 3',11.5)]:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    if name.startswith('Heading'):
        s.paragraph_format.keep_with_next=True
        s.paragraph_format.space_before=Pt(12 if name!='Heading 1' else 0)
        s.paragraph_format.space_after=Pt(6)
    else:
        s.paragraph_format.space_after=Pt(6)
    if name=='Normal': s.paragraph_format.line_spacing=1.08
for p in doc.paragraphs:
    if p.style.name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        for r in p.runs:
            r.font.color.rgb=RGBColor(0,0,0)
            r.font.underline=False
    p.paragraph_format.widow_control=True
    if p.style.name=='Normal': p.paragraph_format.keep_with_next=False
    for border in p._p.xpath('./w:pPr/w:pBdr'): border.getparent().remove(border)

# Bookmark primary sections and create a compact active outline.
bookmark_id=100
outline=[]
for i,p in enumerate(orig):
    if p.style.name=='Heading 1' and ('Продолжение' not in p.text) and not (p.text.startswith('11 ') and i!=163) and not (p.text.startswith('12 ') and i!=205) and i!=1:
        name='section_'+str(i); b=OxmlElement('w:bookmarkStart'); b.set(qn('w:id'),str(bookmark_id)); b.set(qn('w:name'),name)
        e=OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'),str(bookmark_id)); bookmark_id+=1
        p._p.insert(1 if p._p.pPr is not None else 0,b); p._p.append(e); outline.append((p.text,name))
nav=before(orig[16],'Навигация по разделам','Heading 2')
nav.paragraph_format.page_break_before=True
for title,anchor in outline:
    p=before(orig[16]); p.paragraph_format.space_after=Pt(3); link(p,title,anchor=anchor,size=10.5)

def make_table(anchor,rows):
    table=doc.add_table(rows=1, cols=len(rows[0])); table.alignment=WD_TABLE_ALIGNMENT.CENTER
    anchor._p.addprevious(table._tbl)
    for i,value in enumerate(rows[0]): table.rows[0].cells[i].text=value
    for row in rows[1:]:
        cells=table.add_row().cells
        for i,value in enumerate(row): cells[i].text=value
    return table

content=(ROOT/'Docs/Attestation/expanded_sections.md').read_text(encoding='utf-8')
parts=re.split(r'^@@(\d+)\s*$',content,flags=re.M)
for j in range(1,len(parts),2):
    index=int(parts[j]); anchor=orig[index]; lines=parts[j+1].strip().splitlines(); n=0
    while n<len(lines):
        s=lines[n].strip(); n+=1
        if not s: continue
        if s.startswith('### '): before(anchor,s[4:],'Heading 3')
        elif s.startswith('## '): before(anchor,s[3:],'Heading 2')
        elif s.startswith('@ref '):
            path,token,count=[x.strip() for x in s[5:].split('|')]
            p=before(anchor); source_ref(p,path,token,int(count))
        elif s.startswith('@code '):
            path,token,count=[x.strip() for x in s[6:].split('|')]
            f=resolve(path); code_lines=f.read_text(encoding='utf-8-sig').splitlines()
            start=next(i for i,l in enumerate(code_lines) if token in l)
            caption=before(anchor,'Фрагмент '+f.name+' · строки '+str(start+1)+'–'+str(min(len(code_lines),start+int(count))))
            caption.paragraph_format.keep_with_next=True
            for k,line in enumerate(code_lines[start:start+int(count)]):
                p=before(anchor); r=p.add_run(str(start+k+1).rjust(3)+'  '+line)
                r.font.name='Consolas'; r.font.size=Pt(8)
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
                p.paragraph_format.keep_with_next=k<int(count)-1
            before(anchor)
        elif s.startswith('@image '):
            path,caption=[x.strip() for x in s[7:].split('|')]
            p=before(anchor); p.paragraph_format.keep_with_next=True
            picture=p.add_run().add_picture(str(resolve(path)),width=Cm(16.5))
            picture._inline.docPr.set('descr',caption)
            c=before(anchor,caption); c.style=doc.styles['Caption']; c.paragraph_format.space_after=Pt(9)
            for r in c.runs: r.font.color.rgb=RGBColor(0,0,0); r.font.size=Pt(9)
        elif s.startswith('@table '):
            rows=[[v.strip() for v in s[7:].split('|')]]
            while lines[n].strip()!='@endtable': rows.append([v.strip() for v in lines[n].split('|')]); n+=1
            n+=1; make_table(anchor,rows); before(anchor)
        else: before(anchor,s)

# Activate original script references by matching each filename to a verified source file.
game_files=list((ROOT/'Assets/_Game').rglob('*'))
filename_map={}
for f in game_files:
    if f.is_file(): filename_map.setdefault(f.name,[]).append(f)
for p in orig:
    if p.text.startswith('Код проекта:'):
        names=re.findall(r'[\w.-]+\.(?:cs|asmdef|controller)',p.text)
        unique=[]
        for name in names:
            found=filename_map.get(name,[])
            if len(found)==1 and found[0] not in unique: unique.append(found[0])
        if unique:
            q=p.insert_paragraph_before('Файлы для проверки: ')
            for k,f in enumerate(unique):
                if k: q.add_run(' · ')
                link(q,f.name,f.as_uri())
            for r in q.runs: r.font.size=Pt(9)

# Link all original appendix table file names.
for table in doc.tables:
    if table.rows and table.rows[0].cells[0].text=='Файл':
        for row in table.rows[1:]:
            path=row.cells[0].text.strip(); f=ROOT/'Assets/_Game'/path
            if not f.exists(): raise FileNotFoundError(path)
            p=row.cells[0].paragraphs[0]; p.clear(); link(p,path,f.as_uri(),size=9)
    for row in table.rows:
        for cell in row.cells:
            if 'Статус' in cell.text: pass
            if cell.text=='Есть тестовые сценарии; свежего запуска нет':
                cell.text='11 EditMode и 25 PlayMode от 8 октября; контракты повторены 12 октября'
            if cell.text=='Архив не является полным Unity-проектом; часть ссылок ведёт наружу':
                cell.text='Полный проект и зависимости доступны; контроль конфигурации подтверждён'
            if cell.text=='Защищать полный исходный проект, проверить зависимости':
                cell.text='Показать ProjectSettings, Packages и стартовую сцену'
            if cell.text=='Нет нового запуска Editor, компиляции и Test Runner':
                cell.text='Unity-проверки зафиксированы 8 октября; 12 октября повторены исходники и контракты'

# Direct source links in appendix B.
source_docs={286:ROOT,288:Path(r'C:/Users/esinencuP/Downloads/Atestarea_RU.docx'),290:ROOT/'My_Little_Farm_GDD_v0.1.docx',294:ROOT/'My_Little_Farm_Unity_Guide.docx',296:ROOT/'Docs/AI/UnityProjectContext.md',298:ROOT/'Docs/AI/AssetReplacementBacklog.md'}
plan=next(ROOT.glob('*План*.docx')); source_docs[292]=plan
for index,f in source_docs.items():
    if not f.exists(): raise FileNotFoundError(str(f))
    p=orig[index].insert_paragraph_before(); link(p,'Открыть материал проекта',f.as_uri())
for rel in doc.part.rels.values():
    if 'hyperlink' in rel.reltype and 'inputsystem@1.14' in rel.target_ref:
        rel._target=rel.target_ref.replace('inputsystem@1.14','inputsystem@1.20')

# Consistent tables with visible borders, repeating headers and deliberate widths.
for table in doc.tables:
    table.autofit=False
    cols=len(table.columns)
    if cols==6: widths=[3.0,2.6,2.0,2.0,3.0,3.8]
    elif cols==5: widths=[3.5,3.0,4.0,2.7,3.2]
    elif cols==3: widths=[3.4,5.6,7.4]
    elif cols==4: widths=[2.3,2.2,3.6,9.3]
    else: widths=[6.0,10.4]
    if table.rows[0].cells[0].text=='Файл': widths=[8.0,8.4]
    if table.rows[0].cells[0].text=='Категория': widths=[3.0,2.3,11.1]
    if table.rows[0].cells[0].text=='Время': widths=[2.2,2.2,12.0]
    if table.rows[0].cells[0].text=='Риск': widths=[4.1,5.1,7.2]
    for i,w in enumerate(widths): table.columns[i].width=Cm(w)
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+tag); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
    for b in pr.findall(qn('w:tblBorders')): pr.remove(b)
    pr.append(borders)
    for row_index,row in enumerate(table.rows):
        row.height=None
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        if row_index==0:
            rep=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(rep)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF2' if row_index==0 else 'FFFFFF' if row_index%2 else 'F7F9FA'); cp.append(shade)
            margin=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                v=OxmlElement('w:'+side); v.set(qn('w:w'),'85' if side in ['top','bottom'] else '110'); v.set(qn('w:type'),'dxa'); margin.append(v)
            cp.append(margin)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.02
                p.paragraph_format.keep_with_next=(row_index<len(table.rows)-1) if len(table.rows)<=7 else row_index==0
                for r in p.runs:
                    r.font.name='Calibri'; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0,0,0); r.font.bold=row_index==0
                if cols>=5 and ci>=2: p.alignment=WD_ALIGN_PARAGRAPH.CENTER

# Page numbers only; keep the existing restrained header design.
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.clear(); footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('My Little Farm · ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    for r in footer.runs: r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('666666')

doc.core_properties.title='My Little Farm Отчёт по проекту и подготовка к аттестации'
doc.core_properties.subject='Расширенная версия по полному Unity проекту'
doc.core_properties.comments=''
doc.save(OUT)
(WORK/'source_references.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
result=Document(OUT)
words=sum(len(p.text.split()) for p in result.paragraphs)+sum(len(c.text.split()) for t in result.tables for r in t.rows for c in r.cells)
print(json.dumps({'output':str(OUT),'paragraphs':len(result.paragraphs),'tables':len(result.tables),'words':words,'verified_source_references':len(records),'hyperlinks':len([r for r in result.part.rels.values() if 'hyperlink' in r.reltype])},ensure_ascii=False))
